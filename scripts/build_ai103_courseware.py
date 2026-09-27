#!/usr/bin/env python3
"""Build the AI-103 WSQ v5.0 trainer deck, Learner Guide, and Lesson Plan.

The deck is mechanism-led: detailed procedures remain in the Learner Guide and
each instructional slide is anchored by architecture, runtime flow, an exact
contract, editable chart, code, decision rule, failure/control, or evidence.
"""

from __future__ import annotations

import json
import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches as DInches, Pt as DPt, RGBColor as DRGB
from lxml import etree
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn as ppt_qn
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
CW = ROOT / "courseware"
ASSETS = CW / "assets"
LABS = ROOT / "labs"
TITLE = "Developing AI Apps and Agents on Azure (AI-103)"
CODE = "TGS-2023036651"
TSC = "Artificial Intelligence Application in Product Development"
TSC_CODE = "ICT-TEM-4034-1.1"
VERSION = "5.0"
DATE = "27 Sep 2026"
ORG = "Tertiary Infotech Academy Pte Ltd"
UEN = "201200696W"
COURSE_URL = "https://www.tertiarycourses.com.sg/"
REPO_URL = "https://github.com/tertiarycourses/TGS-2023036651-Developing-AI-Apps-and-Agents-on-Azure-AI-103"
LMS_URL = "https://lms-tms.tertiaryinfotech.com/"
AI103_GUIDE = "https://learn.microsoft.com/en-us/credentials/certifications/resources/study-guides/ai-103"
OFFICIAL = json.loads((LABS / "LAB-SELECTION-v5.0.json").read_text())

PPTX = CW / f"{TITLE}-v{VERSION}.pptx"
LG = CW / f"LG-{TITLE}.docx"
LP = CW / f"LP-{TITLE}.docx"
SLIDE_MAP = CW / "slide_map.json"

BLUE = "1F6FEB"
TEAL = "108A73"
VIOLET = "6D3FD2"
AMBER = "C77600"
RED = "C2413A"
GREEN = "16845B"
INK = "161B26"
GREY = "5B6372"
LIGHT = "F5F8FC"
LINE = "D7E0EA"
WHITE = "FFFFFF"
NAVY = "0B1220"
PALETTE = [BLUE, VIOLET, TEAL, GREEN]


def c(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


@dataclass
class Lab:
    num: int
    title: str
    path: Path
    text: str
    objectives: list[str]
    steps: list[str]
    validation: str


def section(text: str, name: str) -> str:
    m = re.search(rf"(?ms)^## {re.escape(name)}\s*$\n(.*?)(?=^## |\Z)", text)
    return m.group(1).strip() if m else ""


def bullets(text: str) -> list[str]:
    return [re.sub(r"^[-*]\s+", "", x.strip()) for x in text.splitlines() if re.match(r"^\s*[-*]\s+", x)]


def load_labs() -> list[Lab]:
    files = sorted(LABS.glob("lab-*.md"))
    if not files:
        files = sorted(LABS.glob("lab-*/README.md"))
    out = []
    for p in files:
        text = p.read_text(encoding="utf-8")
        head = re.search(r"(?m)^# Lab\s+(\d+)\s*[-—:]\s*(.+)$", text)
        if not head:
            continue
        num, title = int(head.group(1)), head.group(2).strip()
        obj = bullets(section(text, "Objectives"))
        step_text = section(text, "Steps")
        steps = [re.sub(r"^\d+\.\s*", "", x).strip() for x in re.findall(r"(?m)^###\s+(.+)$", step_text)]
        validation = " ".join(x.strip() for x in section(text, "Validation").splitlines() if x.strip())
        out.append(Lab(num, title, p, text, obj, steps, validation))
    if len(out) != 10:
        raise RuntimeError(f"Expected 10 labs, found {len(out)}")
    return out


TOPICS = [
    dict(
        title="Plan, secure, and operate Foundry services",
        domain="PLAN & MANAGE · AI-103 25–30%",
        architecture=[("Identity", "Managed identity + least-privilege RBAC"), ("Network", "Private endpoint + controlled egress"), ("Runtime", "Foundry project, models, tools, connections"), ("Evidence", "Azure Monitor logs, metrics, traces, alerts")],
        flow=["Classify workload", "Select service", "Provision identity", "Apply network controls", "Observe SLOs"],
        contracts=[("Project endpoint", "https://<resource>.services.ai.azure.com/api/projects/<project>", "Foundry SDK"), ("Auth", "DefaultAzureCredential", "Entra token; avoid embedded keys"), ("RBAC", "Foundry User", "Project development; check endpoint-specific inference role"), ("Diagnostics", "RequestResponse + Audit", "Send to Log Analytics")],
        code="from azure.identity import DefaultAzureCredential\nfrom azure.ai.projects import AIProjectClient\nclient = AIProjectClient(endpoint=PROJECT_ENDPOINT,\n                         credential=DefaultAzureCredential())\nfor conn in client.connections.list():\n    print(conn.name, conn.type)",
        chart=(['availability', 'p95 latency', 'error rate', 'cost'], [99.95, 82, 1.2, 68], "Indicative dashboard: availability is a percentage; other values are normalized to target. Investigate any measure below its target band."),
        decision=("Does the workload need Foundry-native agents, evaluation, or tracing?", "Use Foundry SDK + project endpoint", "Use the task-specific Foundry Tool or OpenAI-compatible endpoint"),
        failures=[("401 / 403", "Identity lacks data-plane role", "Verify principal, scope, and token audience"), ("429", "Quota or rate limit", "Retry with backoff; request quota or scale units"), ("Timeout", "Network/DNS/private link", "Resolve FQDN from workload subnet"), ("Blind spot", "Diagnostics not routed", "Enable categories and alert ownership")],
        verify="Evidence pack: resource topology, RBAC assignment, private connectivity decision, successful SDK call, and an alert driven by a real metric.",
    ),
    dict(
        title="Responsible AI, content safety, and governance",
        domain="PLAN & MANAGE · AI-103 25–30%",
        architecture=[("Input guard", "Prompt shields + content classification"), ("Model policy", "Deployment filters + approved model/version"), ("Output guard", "Severity thresholds + groundedness checks"), ("Operations", "Human escalation + audit + incident response")],
        flow=["Discover risk", "Measure baseline", "Mitigate controls", "Release with limits", "Monitor & improve"],
        contracts=[("Hate / Sexual / Violence / Self-harm", "severity 0–6", "Block at policy threshold"), ("Prompt attack", "direct / indirect", "Reject, sanitize, or isolate tools"), ("PII", "category + offset + confidence", "Redact or route to reviewer"), ("Decision record", "model, dataset, metric, owner", "Versioned audit evidence")],
        code="POST /contentsafety/text:analyze?api-version=2024-09-01\n{\n  \"text\": user_text,\n  \"categories\": [\"Hate\", \"Violence\", \"SelfHarm\", \"Sexual\"],\n  \"outputType\": \"FourSeverityLevels\"\n}",
        chart=(['allowed', 'review', 'blocked', 'false positive'], [74, 11, 12, 3], "Indicative evaluation set: tune policy against both unsafe misses and legitimate-content false positives; one aggregate safety score is insufficient."),
        decision=("Could a model action create material harm or irreversible change?", "Require human approval and durable audit evidence", "Automate within tested thresholds and monitor exceptions"),
        failures=[("Over-blocking", "Threshold too conservative", "Segment policy by channel and risk"), ("Jailbreak", "Instruction hierarchy compromised", "Isolate untrusted content and tools"), ("Data leak", "Prompt/output contains sensitive data", "Redact, minimize retention, restrict logging"), ("No owner", "Incident cannot be escalated", "Name response owner and SLA")],
        verify="Evidence pack: risk register, test dataset, per-category confusion matrix, filter configuration, exception workflow, and incident rehearsal result.",
        image="ai102-responsible-ai-controls.png",
    ),
    dict(
        title="Generative AI apps and grounded RAG",
        domain="GENERATIVE AI & AGENTS · AI-103 30–35%",
        architecture=[("Ingest", "Chunk, clean, enrich, preserve source IDs"), ("Index", "Text + vector fields in one search index"), ("Retrieve", "Hybrid search + filters + semantic rerank"), ("Generate", "Grounded prompt + citations + evaluation")],
        flow=["Load corpus", "Chunk + embed", "Index metadata", "Retrieve top-k", "Generate with citations"],
        contracts=[("chunk_id", "string, key", "Stable evidence identifier"), ("content_vector", "Collection(Edm.Single)", "Dimensions match embedding model"), ("filter", "tenant_id eq 'T-42'", "Security boundary before generation"), ("top", "5–20", "Recall/latency trade-off")],
        code="response = openai.responses.create(\n    model=DEPLOYMENT,\n    input=[{\"role\": \"user\", \"content\": question}],\n    tools=[{\"type\": \"file_search\",\n            \"vector_store_ids\": [STORE_ID]}]\n)\nprint(response.output_text)",
        chart=(['baseline', 'hybrid', '+ reranker', '+ filters'], [0.54, 0.72, 0.81, 0.79], "Indicative grounded-answer score: hybrid retrieval improves recall; reranking improves ordering; security filters may reduce recall and must be measured per tenant."),
        decision=("Is missing/current knowledge the primary failure?", "Use RAG with governed sources and citations", "Use fine-tuning only for repeatable behavior/style changes"),
        failures=[("Hallucination", "Context absent or ignored", "Require citations; answer 'not found'"), ("Dimension error", "Index and query embeddings differ", "Use one embedding model/config"), ("Leakage", "Filter applied after retrieval", "Filter at query/index boundary"), ("Stale corpus", "Indexer watermark failed", "Track source version and last-success time")],
        verify="Evidence pack: chunk sample, index schema, hybrid query trace, cited answer, groundedness score, latency, and a negative test with no supporting evidence.",
    ),
    dict(
        title="Agents, tools, memory, and multi-agent control",
        domain="GENERATIVE AI & AGENTS · AI-103 30–35%",
        architecture=[("Agent", "Instructions + model + tool policy"), ("Tools", "Typed schema, narrow permissions, timeouts"), ("State", "Conversation/session state separated by user"), ("Control", "Tracing, approvals, budgets, kill switch")],
        flow=["Receive goal", "Plan next action", "Call approved tool", "Inspect observation", "Stop or continue"],
        contracts=[("Tool schema", "name + JSON parameters", "Reject unknown arguments"), ("Identity", "Dedicated agent principal", "No shared admin identity"), ("Budget", "max turns / tokens / cost", "Hard termination condition"), ("Trace", "span, tool call, result, decision", "Replayable evidence")],
        code="agent = project.agents.create_version(\n    agent_name=\"support-triage\",\n    definition=PromptAgentDefinition(\n        model=MODEL_DEPLOYMENT,\n        instructions=SYSTEM_POLICY,\n        tools=[search_tool, ticket_tool]))\nprint(agent.id, agent.version)",
        chart=(['prompt agent', 'hosted agent', 'multi-agent'], [34, 58, 86], "Indicative operational complexity index: add hosting and orchestration only when autonomy, custom code, or specialization justifies extra identity, state, and observability controls."),
        decision=("Can one bounded agent complete the task with one identity and tool set?", "Prefer one agent; simpler state and audit", "Split specialists and add explicit handoff/orchestrator contracts"),
        failures=[("Tool abuse", "Broad permissions or weak schema", "Allow-list actions and validate arguments"), ("Loop", "No stop condition", "Turn/time/cost budgets"), ("Cross-user state", "Memory partition failure", "Key state by tenant + user + session"), ("Opaque handoff", "Agent loses context/owner", "Typed handoff payload + trace link")],
        verify="Evidence pack: agent definition/version, tool schemas, least-privilege identity, approval gate, trace showing tool call/result, budget exhaustion test, and handoff test.",
    ),
    dict(
        title="Image/video generation and multimodal understanding",
        domain="COMPUTER VISION · AI-103 10–15%",
        architecture=[("Input", "Brief or visual media + rights/consent"), ("Model", "Generation, editing, or multimodal analysis"), ("Guardrail", "Safety filter + visual claim review"), ("Output", "Asset, caption, alt text, grounded answer")],
        flow=["Validate media", "Choose model/task", "Generate or analyze", "Check evidence", "Human approve"],
        contracts=[("asset_id", "stable source/output ID", "Provenance and rollback"), ("model_version", "pinned deployment/version", "Reproducible evaluation"), ("claim", "text + source evidence", "Reject unsupported product facts"), ("review", "approved + reason + owner", "Publication gate")],
        code="brief = {\"capacity_ml\": 500, \"allowed_claims\": [\"reusable\"]}\nclaims = [\"reusable\", \"48-hour cooling\"]\nunsupported = set(claims) - set(brief[\"allowed_claims\"])\nassert not unsupported, f\"Reject unsupported claims: {unsupported}\"",
        chart=(['grounded', 'needs review', 'unsafe', 'unsupported'], [62, 20, 8, 10], "Illustrative review set (%): generation quality alone does not grant publication; groundedness, safety, accessibility, and human approval are separate gates."),
        decision=("Does the task create new media or interpret existing visual evidence?", "Use generation/editing with brand and safety review", "Use multimodal understanding with provenance and grounded answers"),
        failures=[("Fabricated claim", "Output exceeds source brief", "Compare with approved attributes"), ("Prompt injection", "Text embedded in image instructs agent", "Treat media as untrusted data"), ("Accessibility", "Alt text omits visual facts", "Evaluate against source image"), ("Rights", "Reference media lacks permission", "Validate provenance and consent")],
        verify="Evidence pack: task selection, source brief, model/version, prompt, generated or sample output, caption and alt text, claim/safety review, and human publication decision.",
    ),
    dict(
        title="Language, speech, translation, and SSML",
        domain="TEXT ANALYSIS · AI-103 10–15%",
        architecture=[("Text", "Language detection, entities, sentiment, PII"), ("Speech", "STT/TTS + custom speech"), ("Translation", "Text/document/speech translation"), ("Experience", "Locale, pronunciation, latency, fallback")],
        flow=["Detect locale", "Transcribe / parse", "Extract meaning", "Translate / synthesize", "Validate quality"],
        contracts=[("language", "en-SG / zh-Hans / auto", "Affects model and voice"), ("offset + length", "Unicode text span", "Preserve source mapping"), ("SSML", "voice, prosody, phoneme, break", "Controlled speech output"), ("recognition result", "text + reason + duration", "Handle NoMatch/Canceled")],
        code="<speak version=\"1.0\" xml:lang=\"en-SG\">\n  <voice name=\"en-SG-LunaNeural\">\n    Your request is <prosody rate=\"-5%\">approved</prosody>.\n  </voice>\n</speak>",
        chart=(['quiet', 'office', 'street', 'call'], [6.8, 11.4, 19.7, 14.2], "Indicative word error rate (%): acoustic context changes recognition quality. Test the real channel and vocabulary; do not accept a studio-only benchmark."),
        decision=("Does the application need deterministic intents/entities or flexible generation?", "Use Azure Language or typed structured extraction", "Use a generative model with schema validation and safety controls"),
        failures=[("NoMatch", "Audio/language mismatch", "Inspect reason and locale"), ("PII leak", "Transcript retained", "Redact before logging"), ("Bad pronunciation", "Voice/phoneme mismatch", "Use SSML phoneme/lexicon"), ("Translation drift", "Domain terms mistranslated", "Glossary/custom model + review")],
        verify="Evidence pack: multilingual samples, PII span/redaction, recognition reasons, SSML output, word-error-rate sample, and human review for domain terms.",
    ),
    dict(
        title="Generative text analysis and translation",
        domain="TEXT ANALYSIS · AI-103 10–15%",
        architecture=[("Input", "Source message + language + provenance"), ("Extract", "Topic, entities, sentiment, summary"), ("Validate", "Typed JSON + exact evidence spans"), ("Route", "PII gate, human review, agent tool")],
        flow=["Detect language", "Extract typed fields", "Verify source spans", "Redact sensitive data", "Route for review"],
        contracts=[("message_id", "stable source ID", "Traceability"), ("evidence_span", "exact source substring", "Reject unsupported claims"), ("sensitive_content", "boolean", "Review and redact"), ("review_required", "boolean", "No unsafe auto-action")],
        code="record = {\"message_id\": \"M1\", \"evidence_span\": \"refund for order 12345\"}\nsource = \"I need a refund for order 12345.\"\nassert record[\"evidence_span\"] in source\n# Validate every generated field before an agent uses it",
        chart=(['grounded extraction', 'translation check', 'PII review', 'human review'], [92, 84, 100, 18], "Illustrative evaluation set (%): report source-grounded extraction, translation fidelity, and sensitive-content handling separately."),
        decision=("Can typed fields be supported by exact source spans?", "Accept schema-valid result after policy checks", "Require human review or a deterministic service fallback"),
        failures=[("Invented entity", "Model guesses a customer fact", "Require exact evidence span"), ("PII leak", "Identifier enters logs or tool", "Redact before routing"), ("Translation drift", "Domain term changes meaning", "Compare source and translation"), ("Bad JSON", "Agent tool receives malformed output", "Validate types and required fields")],
        verify="Evidence pack: synthetic source messages, typed JSON, exact spans, sensitive-content decision, translation comparison, validator output, and agent routing rule.",
    ),
    dict(
        title="Knowledge retrieval, indexing, and grounding",
        domain="INFORMATION EXTRACTION · AI-103 10–15%",
        architecture=[("Source", "Blob/SQL/Cosmos data source"), ("Enrichment", "Indexer + skillset + cache"), ("Index", "Text, filters, vectors, semantic config"), ("Query", "BM25 + vector + RRF + semantic rerank")],
        flow=["Crack documents", "Enrich skills", "Project fields", "Build vector index", "Hybrid query"],
        contracts=[("key", "Edm.String, key=true", "Unique document/chunk ID"), ("vector", "Collection(Edm.Single)", "Dimensions + HNSW profile"), ("search/filter/select", "query contract", "Reduce exposure and payload"), ("scores", "@search.score / rerankerScore", "Different scales; do not compare directly")],
        code="{\n  \"search\": \"maintenance policy\",\n  \"vectorQueries\": [{\"kind\": \"text\", \"text\": \"maintenance policy\",\n+                       \"fields\": \"content_vector\"}],\n  \"queryType\": \"semantic\",\n  \"filter\": \"tenant_id eq 'T-42'\",\n  \"top\": 5\n}",
        chart=(['BM25', 'vector', 'hybrid RRF', '+ semantic'], [0.61, 0.69, 0.78, 0.84], "Indicative nDCG@5: hybrid retrieval often improves coverage, while semantic reranking improves ordering. Validate with your own judged query set."),
        decision=("Must the system support exact terms, filters, and semantic similarity together?", "Use hybrid query + filters + semantic config", "Use the simplest retrieval mode that meets measured relevance") ,
        failures=[("Indexer error", "Skill input path missing", "Inspect enrichment tree and field mappings"), ("Empty vector", "Embedding skill/vectorizer failed", "Check endpoint, dimensions, projection"), ("Poor relevance", "No judged query set", "Create labels and compare retrieval modes"), ("Data leak", "filterable security field absent", "Enforce tenant/ACL filter before results")],
        verify="Evidence pack: index schema, data source/indexer/skillset JSON, successful run history, judged queries, hybrid trace, security-filter negative test, and knowledge-store projection.",
    ),
    dict(
        title="Document Intelligence and Content Understanding",
        domain="INFORMATION EXTRACTION · AI-103 10–15%",
        architecture=[("Input", "Document/image/audio/video + content type"), ("Analyzer", "Prebuilt/custom/composed/Content Understanding"), ("Schema", "Fields, tables, spans, confidence, provenance"), ("Downstream", "Validation, search, workflow, audit")],
        flow=["Classify input", "Select analyzer", "Extract structure", "Validate confidence", "Route downstream"],
        contracts=[("modelId / analyzerId", "prebuilt or versioned custom", "Release identity"), ("pages / spans", "offset + length + polygon", "Trace extracted value to source"), ("confidence", "field-level 0–1", "Drive review threshold"), ("operation-location", "async result URL", "Poll until succeeded/failed")],
        code="poller = client.begin_analyze_document(\n    model_id=\"prebuilt-invoice\",\n    body=document_bytes,\n    content_type=\"application/pdf\")\nresult = poller.result()\nfor doc in result.documents:\n    print(doc.fields[\"InvoiceTotal\"].value_currency)",
        chart=(['invoice', 'receipt', 'contract', 'mixed media'], [0.94, 0.91, 0.79, 0.73], "Indicative field F1: prebuilt models excel on supported formats; custom/composed models or Content Understanding address domain-specific and multimodal schemas."),
        decision=("Is the input a supported document with a stable prebuilt schema?", "Use a prebuilt Document Intelligence model", "Use custom/composed models or Content Understanding for multimodal/domain schemas"),
        failures=[("Wrong model", "Document type misclassified", "Classify first or use composed model"), ("Low confidence", "Layout/domain drift", "Human review + retraining sample"), ("Async timeout", "Polling/size/page limits", "Honor operation status and limits"), ("Lost provenance", "Value stored without span/page", "Persist source URI + region/span")],
        verify="Evidence pack: analyzer/model ID, sample result JSON, source polygons/spans, field-level metrics, review threshold, failure sample, and downstream search/workflow record.",
        image="ai102-multimodal-ingestion.png",
    ),
    dict(
        title="Capstone: production AI solution decision record",
        domain="CAPSTONE · INTEGRATE, OPERATE, AND IMPROVE",
        architecture=[("Experience", "API/UI + accessibility + fallback"), ("Intelligence", "Model/agent + retrieval + extraction"), ("Platform", "Identity, network, data, deployment"), ("Operations", "SLOs, evaluation, audit, incident response")],
        flow=["Frame outcome", "Map risks", "Prototype evidence", "Release gates", "Operate & improve"],
        contracts=[("ADR", "context, decision, alternatives, consequence", "Auditable design rationale"), ("SLO", "availability, p95, quality, cost", "Named owner + error budget"), ("Evaluation", "dataset, metric, threshold, version", "Reproducible release gate"), ("Runbook", "signal, diagnosis, containment, recovery", "Operational readiness")],
        code="release_gate = {\n  \"groundedness\": {\"min\": 0.82},\n  \"unsafe_rate\": {\"max\": 0.01},\n  \"p95_ms\": {\"max\": 2500},\n  \"cost_per_case_sgd\": {\"max\": 0.08}\n}\nassert all(check(metric, rule) for metric, rule in release_gate.items())",
        chart=(['quality', 'safety', 'latency', 'cost'], [86, 99, 78, 84], "Indicative release scorecard normalized to target: a release fails when any hard safety/quality gate fails, even if the average score looks healthy."),
        decision=("Does every material risk have a tested control, owner, and observable signal?", "Approve staged release with rollback", "Return to design/evaluation; do not average away a failed gate"),
        failures=[("Metric gaming", "One aggregate score", "Use per-risk hard gates"), ("No rollback", "Model/index change irreversible", "Blue/green assets + pinned versions"), ("Cost surprise", "Token/vector/skill usage unbounded", "Budgets, caching, quotas, alerts"), ("No evidence", "Decision cannot be audited", "Bundle ADR, tests, traces, runbook")],
        verify="Capstone evidence: architecture and data flow, ADR, threat/risk model, evaluation report, SLO dashboard, cost estimate, deployment/rollback plan, runbook, and improvement backlog.",
        image="ai102-foundry-control-plane.png",
    ),
]

# Keep the five exam domains while following the exact MicrosoftLearning
# exercise order selected for this two-day delivery.
_prior_topics = TOPICS
TOPICS = [dict(_prior_topics[i]) for i in [0, 1, 2, 3, 3, 4, 6, 5, 8, 7]]
for topic, lab in zip(TOPICS, OFFICIAL):
    topic["title"] = lab["title"]
    topic["domain"] = lab["domain"].upper() + " · AI-103"
    topic["verify"] = (
        f"Complete the pinned MicrosoftLearning exercise {lab['source_path']}; "
        "capture its expected output, relevant version/endpoint, one measured result, "
        "and an explained failure or limitation."
    )


class Deck:
    def __init__(self, labs: list[Lab]):
        self.labs = labs
        self.prs = Presentation()
        self.prs.slide_width = Inches(13.333)
        self.prs.slide_height = Inches(7.5)
        self.blank = self.prs.slide_layouts[6]
        self.map: dict[str, int] = {}
        self.kinds: dict[int, str] = {}

    def add(self, key: str, kind: str = "content"):
        s = self.prs.slides.add_slide(self.blank)
        bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, self.prs.slide_width, self.prs.slide_height)
        bg.fill.solid(); bg.fill.fore_color.rgb = c(WHITE); bg.line.fill.background()
        self.map[key] = len(self.prs.slides)
        self.kinds[len(self.prs.slides)] = kind
        return s

    def shape(self, s, typ, x, y, w, h, fill=LIGHT, line=LINE, radius=False):
        sh = s.shapes.add_shape(typ, Inches(x), Inches(y), Inches(w), Inches(h))
        sh.fill.solid(); sh.fill.fore_color.rgb = c(fill)
        if line: sh.line.color.rgb = c(line); sh.line.width = Pt(0.8)
        else: sh.line.fill.background()
        return sh

    def text(self, s, x, y, w, h, value, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT, font="Arial", valign=MSO_ANCHOR.TOP):
        box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame; tf.clear(); tf.word_wrap = True
        tf.margin_left = tf.margin_right = Pt(0); tf.margin_top = tf.margin_bottom = Pt(0)
        tf.vertical_anchor = valign
        p = tf.paragraphs[0]; p.alignment = align
        r = p.add_run(); r.text = str(value); r.font.name = font; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = c(color)
        return box

    def header(self, s, title, kicker, accent=BLUE):
        self.shape(s, MSO_SHAPE.RECTANGLE, 0, 0, .22, 1.48, accent, None)
        self.text(s, .72, .34, 11.85, .22, kicker.upper(), 12, accent, True)
        size = 29 if len(title) <= 48 else 24
        self.text(s, .72, .68, 11.85, .58, title, size, INK, True)
        self.shape(s, MSO_SHAPE.RECTANGLE, .72, 1.50, 11.85, .01, LINE, None)

    def footer(self, s):
        n = len(self.prs.slides)
        self.text(s, .72, 7.02, 5.3, .18, f"{TITLE} · {CODE}", 8, GREY)
        self.text(s, 5.1, 7.02, 4.2, .18, f"© 2026 {ORG}", 8, GREY, False, PP_ALIGN.CENTER)
        self.text(s, 11.95, 7.02, .6, .18, n, 8, GREY, True, PP_ALIGN.RIGHT)

    def cover(self):
        s = self.add("cover", "admin")
        img = ASSETS / "ai102-foundry-control-plane.png"
        if img.exists(): s.shapes.add_picture(str(img), Inches(6.25), Inches(.55), width=Inches(6.5), height=Inches(5.8))
        self.shape(s, MSO_SHAPE.RECTANGLE, 0, 0, .24, 7.5, BLUE, None)
        logo = ASSETS / "tertiary-logo.png"
        if logo.exists(): s.shapes.add_picture(str(logo), Inches(.72), Inches(.52), height=Inches(.72))
        self.text(s, .72, 1.62, 5.55, .28, "WSQ · INSTRUCTOR-LED · 2 DAYS", 13, BLUE, True)
        self.text(s, .72, 2.05, 5.6, 1.65, TITLE, 32, INK, True)
        self.text(s, .72, 4.02, 5.55, .34, f"Course Code: {CODE}", 15, GREY)
        self.text(s, .72, 4.46, 5.55, .56, f"{ORG}\nUEN {UEN}", 12.5, GREY)
        self.shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, .72, 5.42, 3.05, .64, "EAF3FF", "BBD7F2")
        self.text(s, .96, 5.62, 2.6, .22, f"Version {VERSION} · {DATE}", 13, BLUE, True)
        self.footer(s)

    def section_slide(self, key, number, title, subtitle, accent):
        s = self.add(key, "divider")
        self.shape(s, MSO_SHAPE.RECTANGLE, .72, 1.90, 11.85, 3.60, LIGHT, None)
        self.shape(s, MSO_SHAPE.RECTANGLE, .72, 1.90, .14, 3.60, accent, None)
        self.text(s, 1.40, 2.35, 9.8, .32, f"TOPIC {number:02d}", 15, accent, True)
        self.text(s, 1.40, 2.85, 9.9, 1.1, title, 34 if len(title)<48 else 28, INK, True)
        self.text(s, 1.40, 4.22, 9.8, .68, subtitle, 15, GREY)
        self.text(s, 11.0, 2.08, 1.0, .9, f"{number:02d}", 54, "DDE6F0", True, PP_ALIGN.RIGHT)
        self.footer(s)

    def cards(self, key, title, kicker, items, accent=BLUE):
        s = self.add(key); self.header(s, title, kicker, accent)
        xs = [.72, 3.72, 6.72, 9.72]
        for i, ((head, body), x) in enumerate(zip(items, xs), 1):
            col = PALETTE[(i-1)%4]
            self.shape(s, MSO_SHAPE.RECTANGLE, x, 1.92, 2.72, 2.55, LIGHT, None)
            self.shape(s, MSO_SHAPE.RECTANGLE, x, 1.92, .09, 2.55, col, None)
            self.shape(s, MSO_SHAPE.OVAL, x+.24, 2.12, .46, .46, col, None)
            self.text(s, x+.24, 2.23, .46, .18, i, 13, WHITE, True, PP_ALIGN.CENTER)
            self.text(s, x+.82, 2.08, 1.62, .44, head, 13.5, INK, True)
            self.text(s, x+.28, 2.73, 2.18, 1.25, body, 11.5, INK)
        self.shape(s, MSO_SHAPE.RECTANGLE, .72, 4.72, 11.85, 1.35, LIGHT, None)
        self.text(s, 1.0, 4.94, 11.2, .2, "ARCHITECTURE RULE", 12, accent, True)
        self.text(s, 1.0, 5.34, 11.1, .48, "Keep identity, data, runtime state, and evidence boundaries explicit; every connection needs an owner, contract, and observable failure signal.", 13, INK)
        self.footer(s)

    def process(self, key, title, kicker, labels, accent=BLUE):
        s = self.add(key); self.header(s, title, kicker, accent)
        xs = [.85, 3.24, 5.64, 8.03, 10.43]
        for i, (x, label) in enumerate(zip(xs, labels), 1):
            self.shape(s, MSO_SHAPE.RECTANGLE, x, 2.10, 2.05, 3.15, LIGHT, LINE)
            self.shape(s, MSO_SHAPE.RECTANGLE, x, 2.10, 2.05, .10, PALETTE[(i-1)%4], None)
            self.shape(s, MSO_SHAPE.OVAL, x+.62, 2.52, .82, .82, PALETTE[(i-1)%4], None)
            self.text(s, x+.62, 2.72, .82, .34, i, 28, WHITE, True, PP_ALIGN.CENTER)
            self.text(s, x+.16, 3.68, 1.73, 1.16, label, 14, INK, True, PP_ALIGN.CENTER)
            if i < 5:
                conn = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x+2.05), Inches(3.67), Inches(xs[i]), Inches(3.67))
                conn.line.color.rgb = c(accent); conn.line.width = Pt(2)
                ln = conn._element.spPr.ln
                end = etree.SubElement(ln, ppt_qn("a:tailEnd")); end.set("type", "triangle")
        self.shape(s, MSO_SHAPE.RECTANGLE, .85, 5.56, 11.63, .72, "EAF3FF", None)
        self.text(s, 1.08, 5.80, 11.15, .26, "Each stage emits evidence for the next stage; a missing transition or identifier is an observable contract failure.", 12.5, BLUE, True, PP_ALIGN.CENTER)
        self.footer(s)

    def table(self, key, title, kicker, rows, accent=BLUE):
        s = self.add(key); self.header(s, title, kicker, accent)
        data = [["Contract", "Exact value / field", "Engineering consequence"]] + [list(r) for r in rows]
        tab = s.shapes.add_table(len(data), 3, Inches(.72), Inches(1.90), Inches(11.85), Inches(4.45)).table
        tab.columns[0].width=Inches(2.35); tab.columns[1].width=Inches(4.45); tab.columns[2].width=Inches(5.05)
        for ri,row in enumerate(data):
            for ci,val in enumerate(row):
                cell=tab.cell(ri,ci); cell.text=str(val); cell.margin_left=cell.margin_right=Inches(.12); cell.margin_top=cell.margin_bottom=Inches(.08)
                cell.fill.solid(); cell.fill.fore_color.rgb=c(accent if ri==0 else (LIGHT if ri%2 else WHITE))
                for p in cell.text_frame.paragraphs:
                    p.font.name="Arial"; p.font.size=Pt(12 if ri else 12.5); p.font.bold=(ri==0 or ci==0); p.font.color.rgb=c(WHITE if ri==0 else INK)
        self.footer(s)

    def code_points(self, key, title, kicker, code, accent=BLUE):
        s=self.add(key); self.header(s,title,kicker,accent)
        self.shape(s,MSO_SHAPE.ROUNDED_RECTANGLE,.72,1.90,7.15,4.75,NAVY,None)
        self.shape(s,MSO_SHAPE.RECTANGLE,.72,1.90,.10,4.75,accent,None)
        lines=code.splitlines(); fs=14 if len(lines)<=9 else 11.5
        self.text(s,1.02,2.20,6.52,4.05,code,fs,"DCEBFF",False,PP_ALIGN.LEFT,"Consolas")
        points=[("AUTH", "Prefer Entra identity; keep secrets out of source."), ("CONTRACT", "Pin endpoint/API/model versions and validate schema."), ("EVIDENCE", "Log request IDs, deployment versions, and outcomes."), ("FAILURE", "Handle auth, quota, timeout, and invalid-state paths.")]
        y=1.90
        for i,(h,b) in enumerate(points):
            col=PALETTE[i]
            self.shape(s,MSO_SHAPE.RECTANGLE,8.15,y,4.42,1.02,LIGHT,None)
            self.shape(s,MSO_SHAPE.RECTANGLE,8.15,y,.09,1.02,col,None)
            self.text(s,8.45,y+.16,3.7,.18,h,11,col,True)
            self.text(s,8.45,y+.46,3.72,.42,b,11.2,INK)
            y+=1.18
        self.footer(s)

    def chart(self,key,title,kicker,categories,values,insight,accent=BLUE):
        s=self.add(key); self.header(s,title,kicker,accent)
        cd=ChartData(); cd.categories=categories; cd.add_series("Indicative teaching values",values)
        chart=s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Inches(.88),Inches(1.92),Inches(11.55),Inches(3.65),cd).chart
        chart.has_title=False; chart.has_legend=False
        chart.value_axis.tick_labels.font.name="Arial"; chart.value_axis.tick_labels.font.size=Pt(11)
        chart.category_axis.tick_labels.font.name="Arial"; chart.category_axis.tick_labels.font.size=Pt(11)
        ser=chart.series[0]; ser.format.fill.solid(); ser.format.fill.fore_color.rgb=c(accent)
        plot=chart.plots[0]; plot.has_data_labels=True; plot.data_labels.position=XL_LABEL_POSITION.OUTSIDE_END; plot.data_labels.font.name="Arial"; plot.data_labels.font.size=Pt(11); plot.data_labels.number_format="0.00"
        self.shape(s,MSO_SHAPE.RECTANGLE,.88,5.78,11.55,.76,LIGHT,None)
        self.text(s,1.12,5.94,2.05,.18,"WHAT THE DATA SHOWS",11,accent,True)
        self.text(s,3.14,5.90,8.95,.46,insight,11.4,INK)
        self.footer(s)

    def decision(self,key,title,kicker,question,yes,no,accent=BLUE):
        s=self.add(key); self.header(s,title,kicker,accent)
        self.shape(s,MSO_SHAPE.DIAMOND,4.78,1.88,3.78,2.60,"EAF3FF",accent)
        self.text(s,5.46,2.63,2.42,.82,question,13,INK,True,PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        self.shape(s,MSO_SHAPE.ROUNDED_RECTANGLE,.95,4.85,5.3,1.25,"E8F7EE",GREEN)
        self.shape(s,MSO_SHAPE.ROUNDED_RECTANGLE,7.08,4.85,5.3,1.25,"FFF7E6",AMBER)
        self.text(s,1.25,5.08,.65,.22,"YES",12,GREEN,True)
        self.text(s,1.95,5.02,3.95,.52,yes,13,INK,True)
        self.text(s,7.38,5.08,.55,.22,"NO",12,AMBER,True)
        self.text(s,8.00,5.02,4.0,.52,no,13,INK,True)
        for x1,x2,col in [(5.38,3.62,GREEN),(7.96,9.72,AMBER)]:
            conn=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(4.08),Inches(x2),Inches(4.85)); conn.line.color.rgb=c(col); conn.line.width=Pt(2)
            end=etree.SubElement(conn._element.spPr.ln,ppt_qn("a:tailEnd")); end.set("type","triangle")
        self.footer(s)

    def failures(self,key,title,kicker,items,accent=RED):
        s=self.add(key); self.header(s,title,kicker,accent)
        xs=[.72,3.72,6.72,9.72]
        for i,((sym,cause,control),x) in enumerate(zip(items,xs),1):
            self.shape(s,MSO_SHAPE.RECTANGLE,x,1.92,2.72,3.72,WHITE,LINE)
            self.shape(s,MSO_SHAPE.RECTANGLE,x,1.92,2.72,.12,RED,None)
            self.shape(s,MSO_SHAPE.OVAL,x+.22,2.26,.48,.48,RED,None)
            self.text(s,x+.22,2.37,.48,.18,"!",15,WHITE,True,PP_ALIGN.CENTER)
            self.text(s,x+.82,2.26,1.56,.42,sym,13,INK,True)
            self.text(s,x+.26,3.00,2.18,.18,"LIKELY CAUSE",10,RED,True)
            self.text(s,x+.26,3.30,2.18,.72,cause,11.2,INK)
            self.text(s,x+.26,4.25,2.18,.18,"CONTROL / RECOVERY",10,GREEN,True)
            self.text(s,x+.26,4.55,2.18,.72,control,11.2,INK)
        self.shape(s,MSO_SHAPE.RECTANGLE,.72,5.92,11.85,.52,"FEF2F2",None)
        self.text(s,1.0,6.08,11.3,.18,"A control is incomplete until it has a detection signal, owner, containment action, and recovery test.",12,RED,True,PP_ALIGN.CENTER)
        self.footer(s)

    def activity(self,key,lab:Lab,topic,accent=TEAL):
        s=self.add(key); self.header(s,f"Lab {lab.num:02d}: {lab.title}","HANDS-ON PRACTICE BRIDGE",accent)
        self.text(s,.82,1.80,11.65,.54,"Mechanism focus: "+topic["title"],14,INK,True)
        names=lab.steps[:5] or ["Frame", "Configure", "Execute", "Validate", "Record"]
        while len(names)<5: names.append(["Frame","Configure","Execute","Validate","Record"][len(names)])
        short=[]
        for n in names[:5]:
            words=n.split(); short.append(" ".join(words[:5]))
        xs=[.85,3.24,5.64,8.03,10.43]
        for i,(x,label) in enumerate(zip(xs,short),1):
            self.shape(s,MSO_SHAPE.ROUNDED_RECTANGLE,x,2.62,2.05,1.75,LIGHT,LINE)
            self.shape(s,MSO_SHAPE.OVAL,x+.70,2.88,.65,.65,PALETTE[(i-1)%4],None)
            self.text(s,x+.70,3.04,.65,.24,i,19,WHITE,True,PP_ALIGN.CENTER)
            self.text(s,x+.18,3.70,1.69,.48,label,11.2,INK,True,PP_ALIGN.CENTER)
            if i<5: self.text(s,x+2.08,3.28,.25,.34,"▶",15,accent,True,PP_ALIGN.CENTER)
        self.shape(s,MSO_SHAPE.RECTANGLE,.85,4.76,11.63,.82,"E8F7EE",None)
        self.text(s,1.12,4.96,2.05,.18,"YOU WILL PRODUCE",11,GREEN,True)
        self.text(s,3.08,4.90,8.95,.42,lab.validation or topic["verify"],12,INK,True)
        self.text(s,.85,5.92,11.63,.34,f"Complete click-by-click procedure: Learner Guide · Lab {lab.num:02d} · {lab.path.relative_to(ROOT)}",11,GREY)
        self.footer(s)

    def official_source_slides(self, n: int, topic: dict, accent: str):
        """Eight technical teaching views anchored to the pinned, unchanged lab."""
        meta = OFFICIAL[n - 1]
        repo = LABS / "official" / meta["source_repository"]
        source = LABS / meta["source_path"]
        body = source.read_text(encoding="utf-8")
        manifest = json.loads((LABS / "OFFICIAL-SOURCE-MANIFEST.json").read_text())
        sha = manifest[meta["source_repository"]]["commit"][:12]
        paths = list(manifest[meta["source_repository"]]["files"])
        referenced = [x.rstrip(".,);`'") for x in re.findall(r"(?:Labfiles|labfiles)/[\w./-]+", body)
                      if (repo / x.rstrip(".,);`'")).exists()]
        source_rel = str(source.relative_to(repo))
        extras = [source_rel, "LICENSE"] + [p for p in ("README.md", "readme.md", "index.md") if (repo / p).exists()]
        referenced = list(dict.fromkeys(referenced + extras))
        for p in paths:
            if len(referenced) >= 4:
                break
            if p not in referenced:
                referenced.append(p)
        fragments = [x.strip() for x in re.findall(r"```[^\n]*\n(.*?)```", body, re.S)
                     if len(x.strip()) >= 80 and re.search(r"\b(import|from|az|pip|curl|client|response|model)\b|[={}]", x)]
        fragments.sort(key=len, reverse=True)
        asset_prefixes = tuple(p.rsplit("/", 1)[0] + "/" for p in referenced if p.startswith(("Labfiles/", "labfiles/")))
        for p in paths:
            if len(fragments) >= 3:
                break
            if p.endswith(".py") and asset_prefixes and p.startswith(asset_prefixes) and not any(v in p.lower() for v in ("solution/", "test_")):
                lines = (repo / p).read_text(encoding="utf-8", errors="ignore").splitlines()
                fragment = "\n".join(x for x in lines if x.strip() and not x.lstrip().startswith("#"))[:650]
                if fragment:
                    fragments.append(fragment)
        for fallback in [topic["code"],
                         "\n".join(f"{a}: {b}  # {c}" for a, b, c in topic["contracts"]),
                         "\n".join(f"{a}: {b}" for a, b in topic["architecture"])]:
            if len(fragments) >= 3:
                break
            if fallback not in fragments:
                fragments.append(fallback)
        fragments = ["\n".join(x.splitlines()[:10])[:680] for x in fragments[:3]]
        headings = [re.sub(r"[#*`]", "", x).strip() for x in re.findall(r"(?m)^##+ (.+)$", body)]
        headings = [x for x in headings if x.lower() not in ("clean up", "cleanup", "summary")]
        phases = headings[:5]
        while len(phases) < 5:
            phases.append(["Configure", "Run", "Inspect", "Evaluate", "Record"][len(phases)])
        self.table(f"topic{n:02d}_source", f"Official lab {n:02d} · pinned source", "MICROSOFTLEARNING · EXACT FILES",
                   [("Repository", meta["source_repository"], "MIT source snapshot"),
                    ("Revision", sha, "SHA-256 per file in source manifest"),
                    ("Exercise", source.name, "Unchanged instruction file"),
                    ("Domain", meta["domain"], "AI-103 study guide")], accent)
        self.cards(f"topic{n:02d}_inputs", f"Official lab {n:02d} · input boundaries", "ASSET AND RUNTIME CONTRACT",
                   [("EXERCISE", source.name[:42]), ("ASSET 1", referenced[0][-46:]),
                    ("ASSET 2", referenced[1][-46:]), ("RUNTIME", "Azure subscription · Python 3.13 · Foundry")], accent)
        for j, fragment in enumerate(fragments, 1):
            self.code_points(f"topic{n:02d}_sourcecode{j}", f"Official lab {n:02d} · implementation {j}",
                             "SOURCE EXCERPT OR DOMAIN CONTRACT · SEE UNCHANGED LAB", fragment, accent)
        self.process(f"topic{n:02d}_sourceflow", f"Official lab {n:02d} · mechanism phases",
                     "SOURCE EXERCISE · NAMED RUNTIME STATES", phases[:5], accent)
        self.table(f"topic{n:02d}_assets", f"Official lab {n:02d} · source assets", "REPRODUCIBLE FILE CONTRACT",
                   [(f"Asset {j}", p[:65], "Exact path in pinned source") for j, p in enumerate(referenced[:4], 1)], accent)
        self.table(f"topic{n:02d}_evidence", f"Official lab {n:02d} · evidence record", "INPUT · OUTPUT · CONTROL · DECISION",
                   [("Input", "resource / model / source version", "Capture before run"),
                    ("Output", "result / request ID / trace", "Compare to exercise expectation"),
                    ("Metric", "quality / latency / cost", "State unit and sample"),
                    ("Control", "failure / limit / next action", "Explain product implication")], accent)

    def verify(self,key,lab:Lab,topic,accent=GREEN):
        s=self.add(key); self.header(s,f"Lab {lab.num:02d}: acceptance evidence","VERIFY · DIAGNOSE · RECORD",accent)
        self.shape(s,MSO_SHAPE.ROUNDED_RECTANGLE,.85,1.95,7.25,3.78,"E8F7EE",GREEN)
        self.text(s,1.18,2.28,6.55,.25,"PASS WHEN",13,GREEN,True)
        self.text(s,1.18,2.78,6.55,2.15,topic["verify"],16,INK,True)
        self.shape(s,MSO_SHAPE.ROUNDED_RECTANGLE,8.35,1.95,4.12,3.78,"FEF2F2",RED)
        self.text(s,8.68,2.28,3.5,.25,"IF IT FAILS",13,RED,True)
        self.text(s,8.68,2.78,3.35,2.2,"1. Capture request/operation ID\n2. Verify identity and endpoint\n3. Inspect schema/model version\n4. Check quota/network/status\n5. Preserve evidence before retry",13,INK)
        self.shape(s,MSO_SHAPE.RECTANGLE,.85,5.98,11.62,.48,LIGHT,None)
        self.text(s,1.08,6.12,11.15,.18,"Assessment bridge: submit the observable artifact and explain why the evidence proves the stated TSC ability.",11.4,BLUE,True,PP_ALIGN.CENTER)
        self.footer(s)

    def image_slide(self,key,title,kicker,image_name,labels,accent=BLUE):
        p=ASSETS/image_name
        if not p.exists(): return
        s=self.add(key); self.header(s,title,kicker,accent)
        s.shapes.add_picture(str(p),Inches(.72),Inches(1.82),width=Inches(8.1),height=Inches(4.85))
        y=1.92
        for i,(h,b) in enumerate(labels):
            col=PALETTE[i%4]
            self.shape(s,MSO_SHAPE.RECTANGLE,9.05,y,3.52,1.02,LIGHT,None)
            self.shape(s,MSO_SHAPE.RECTANGLE,9.05,y,.09,1.02,col,None)
            self.text(s,9.32,y+.14,2.95,.18,h,11,col,True)
            self.text(s,9.32,y+.43,2.95,.45,b,10.8,INK)
            y+=1.16
        self.footer(s)

    def admin_cards(self,key,title,kicker,items,accent=BLUE):
        self.cards(key,title,kicker,items,accent)

    def trainer(self,key,blank=False):
        s=self.add(key,"admin"); self.header(s,"About the Trainer","TRAINING TEAM PROFILE" if blank else "NAMED TRAINER",BLUE)
        col=TEAL if blank else BLUE
        self.shape(s,MSO_SHAPE.RECTANGLE,.85,1.92,3.62,4.45,LIGHT,None)
        self.shape(s,MSO_SHAPE.OVAL,1.88,2.45,1.56,1.56,col,None)
        self.text(s,1.88,2.88,1.56,.48,"TIA" if blank else "AA",25 if blank else 30,WHITE,True,PP_ALIGN.CENTER)
        self.text(s,1.18,4.30,2.96,.60,"Tertiary Infotech\nAcademy Trainer" if blank else "Dr. Alfred Ang",18,INK,True,PP_ALIGN.CENTER)
        self.text(s,1.18,5.03,2.96,.68,"Microsoft Azure AI\ntraining team" if blank else "Principal Trainer\nAI, automation, and software engineering",11.5,GREY,False,PP_ALIGN.CENTER)
        rows=[("NAME","Tertiary Infotech Academy Trainer" if blank else "Dr. Alfred Ang"),("TITLE / DESIGNATION","WSQ Technical Trainer" if blank else "Principal Trainer"),("QUALIFICATIONS","Azure AI and adult-learning credentials" if blank else "PhD; Azure AI and software engineering"),("AREAS OF EXPERTISE","Microsoft Foundry, Azure AI services, MLOps" if blank else "AI solutions, agents, automation, apps"),("EXPERIENCE","Applied cloud AI delivery and facilitation" if blank else "Industry implementation and adult learning"),("CONTACT","Introduced during the opening briefing" if blank else "Shared during class")]
        y=1.92
        for i,(h,b) in enumerate(rows):
            self.shape(s,MSO_SHAPE.RECTANGLE,4.82,y,7.75,.63,LIGHT,None); self.shape(s,MSO_SHAPE.RECTANGLE,4.82,y,.08,.63,PALETTE[i%4],None)
            self.text(s,5.12,y+.14,2.2,.18,h,10,PALETTE[i%4],True); self.text(s,7.35,y+.12,4.78,.28,b,11.4,INK)
            y+=.74
        self.footer(s)

    def assessment_flow(self,key,front=False):
        self.process(key,"Assessment Flow","BRIEFING" if front else "END-OF-COURSE REMINDER",["TRAQOM digital attendance","Assessment digital attendance","WA-SAQ · 60 min","PP · 60 min","Submit on LMS + sign record"],VIOLET)

    def hyperlink_card(self,key,title,url,body,accent=BLUE,display_text=None):
        s=self.add(key,"admin"); self.header(s,title,"CLICKABLE RESOURCE",accent)
        self.shape(s,MSO_SHAPE.ROUNDED_RECTANGLE,.92,2.00,11.48,3.62,LIGHT,LINE)
        self.text(s,1.28,2.48,10.7,.44,body,16,INK,True,PP_ALIGN.CENTER)
        box=self.text(s,1.55,3.38,9.95,.52,display_text or url,15,accent,True,PP_ALIGN.CENTER)
        box.text_frame.paragraphs[0].runs[0].hyperlink.address=url
        self.shape(s,MSO_SHAPE.ROUNDED_RECTANGLE,4.76,4.45,3.82,.62,accent,None)
        self.text(s,4.76,4.64,3.82,.22,"OPEN RESOURCE",13,WHITE,True,PP_ALIGN.CENTER)
        self.footer(s)

    def build(self):
        self.cover()
        self.section_slide("admin_start",0,"Course administration","WSQ requirements · learning outcomes · evidence expectations",BLUE)
        self.admin_cards("admin_attendance","Digital Attendance (Mandatory)","TRAQOM · SSG",[("AM","Complete the morning attendance window"),("PM","Complete the afternoon attendance window"),("ASSESSMENT","Complete assessment attendance before sitting papers"),("ELIGIBILITY","At least 75% attendance and Competent outcome")],TEAL)
        self.trainer("admin_trainer_general",True); self.trainer("admin_trainer_named",False)
        self.admin_cards("admin_ground_rules","Ground Rules","LEARNING ENVIRONMENT",[("PUNCTUAL","Return from breaks on time"),("PARTICIPATE","Explain decisions and evidence"),("PROTECT DATA","Use synthetic/non-confidential inputs"),("ASK EARLY","Raise blockers before they cascade")],BLUE)
        self.admin_cards("admin_outcomes","Learning Outcomes","TSC ALIGNMENT",[("LO1","Analyse Azure AI algorithms and efficiency"),("LO2","Evaluate strengths and limitations"),("LO3","Assess feasibility and improvements"),("EVIDENCE","Explain what proves each conclusion")],VIOLET)
        self.admin_cards("admin_schedule","Two-Day Lesson Plan","9:00 AM–6:00 PM",[("DAY 1 AM","Foundry project, model evaluation, chat app"),("DAY 1 PM","Agents, custom tools, vision"),("DAY 2 AM","Text analysis, text agent, content understanding"),("DAY 2 PM","Knowledge mining + 2-hour assessment")],TEAL)
        self.admin_cards("admin_exam_status","AI-103 Course Scope","CURRENT AS OF 27 SEP 2026",[("EXAM","AI-103: Developing AI Apps and Agents on Azure"),("SCOPE","Skills measured from 16 Apr 2026"),("COURSE","WSQ TGS-2023036651 · two days"),("FOCUS","Foundry apps, agents, multimodal and extraction")],AMBER)
        self.hyperlink_card("admin_ai103","Current Microsoft certification pathway",AI103_GUIDE,"AI-103: Developing AI Apps and Agents on Azure · skills measured from 16 Apr 2026",BLUE,"Microsoft Learn · AI-103 Study Guide")
        self.admin_cards("admin_briefing","Briefing for Assessment","READ BEFORE ASSESSMENT",[("OPEN BOOK","Use approved slides and Learner Guide"),("INDIVIDUAL","No discussion or shared answers"),("EVIDENCE","Answer every K/A criterion"),("SUBMISSION","Upload through the LMS")],VIOLET)
        self.admin_cards("admin_assessment","Assessment","APPROVED PLAN",[("WA-SAQ","6 open-ended questions · K1–K6"),("PP","6 lab-evidence tasks · A1–A6"),("TIMING","60 minutes each"),("OUTCOME","Competent / Not Yet Competent")],VIOLET)
        self.assessment_flow("admin_assessment_flow",True)
        self.hyperlink_card("admin_lms","Courseware and Assessment on the LMS",LMS_URL,"Download the learner materials and submit the two candidate papers on the LMS",TEAL,"LMS-TMS · Courseware and Assessment")
        self.hyperlink_card("admin_labs","Access the Hands-On Labs",REPO_URL,"Open the ten numbered lab guides and the pinned, unchanged MicrosoftLearning source snapshots",BLUE,"GitHub · AI-103 Official Lab Materials")
        self.section_slide("day1",0,"Day 1","Foundry planning · models · chat · agents · vision",BLUE)
        for i,(topic,lab) in enumerate(zip(TOPICS,self.labs),1):
            if i==7: self.section_slide("day2",0,"Day 2","Text analysis · agents · content understanding · knowledge mining",TEAL)
            accent=PALETTE[(i-1)%len(PALETTE)]
            self.section_slide(f"topic{i:02d}",i,topic["title"],topic["domain"],accent)
            if topic.get("image"):
                self.image_slide(f"topic{i:02d}_image",topic["title"],"GENERATED CONCEPT VISUAL + EDITABLE LABELS",topic["image"],topic["architecture"],accent)
            self.cards(f"topic{i:02d}_arch",topic["title"]+" · architecture",topic["domain"],topic["architecture"],accent)
            self.process(f"topic{i:02d}_flow",topic["title"]+" · runtime trace","NAMED STATES AND TRANSITIONS",topic["flow"],accent)
            self.table(f"topic{i:02d}_contract",topic["title"]+" · exact contract","FIELDS · ENDPOINTS · PARAMETERS",topic["contracts"],accent)
            self.code_points(f"topic{i:02d}_code",topic["title"]+" · implementation excerpt","RUNNABLE / CONFIGURABLE ANCHOR",topic["code"],accent)
            cats,vals,insight=topic["chart"]
            self.chart(f"topic{i:02d}_chart",topic["title"]+" · engineering metric","EDITABLE NATIVE POWERPOINT CHART",cats,vals,insight,accent)
            q,yes,no=topic["decision"]
            self.decision(f"topic{i:02d}_decision",topic["title"]+" · decision rule","CHOOSE FROM EVIDENCE",q,yes,no,accent)
            self.failures(f"topic{i:02d}_failure",topic["title"]+" · failure and control","DIAGNOSE BEFORE RETRY",topic["failures"])
            self.activity(f"lab{i:02d}_activity",lab,topic,accent)
            self.official_source_slides(i, topic, accent)
            self.verify(f"lab{i:02d}_verify",lab,topic)
        self.section_slide("closing",0,"Assessment and course close","Evidence submission · assessment flow · support",VIOLET)
        self.hyperlink_card("closing_support","Course enquiries and support",COURSE_URL,"Contact Tertiary Courses for current course details, schedules, and registration",BLUE,"Tertiary Courses · Course Enquiries")
        self.admin_cards("closing_assessment","Assessment","FINAL REMINDER",[("WA-SAQ","60 min · K1–K6"),("PP","60 min · A1–A6"),("OPEN BOOK","Approved materials only"),("SUBMIT","Candidate papers on LMS")],VIOLET)
        self.assessment_flow("closing_flow")
        self.admin_cards("closing_attendance","Digital Attendance (Mandatory)","TRAQOM · SSG",[("SCAN","Use the LMS/TMS QR code"),("ASSESSMENT","Complete digital attendance"),("VERIFY","Confirm attendance recorded"),("SIGN","Complete Assessment Summary Record")],TEAL)
        self.section_slide("closing_thanks",0,"Thank You","You can now explain, test, and operate an end-to-end Azure AI solution",BLUE)
        self.apply_transitions()
        self.prs.save(PPTX)
        SLIDE_MAP.write_text(json.dumps(self.map,indent=2),encoding="utf-8")
        return self.map

    def apply_transitions(self):
        for n,s in enumerate(self.prs.slides,1):
            sld=s._element
            for old in sld.findall(ppt_qn("p:transition")): sld.remove(old)
            tr=etree.SubElement(sld,ppt_qn("p:transition")); tr.set("spd","med" if self.kinds.get(n)=="divider" else "fast"); tr.set("advClick","1")
            if self.kinds.get(n)=="divider": etree.SubElement(tr,ppt_qn("p:push")).set("dir","l")
            else: etree.SubElement(tr,ppt_qn("p:fade"))
            sld.append(tr)


def drgb(value: str) -> DRGB:
    return DRGB.from_string(value)


def doc_font(run,size=11,bold=False,color=INK,font="Arial"):
    run.font.name=font; run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"),font); run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"),font)
    run.font.size=DPt(size); run.bold=bold; run.font.color.rgb=drgb(color)


def setup_doc(doc:Document):
    sec=doc.sections[0]; sec.page_width=DInches(8.27); sec.page_height=DInches(11.69); sec.top_margin=sec.bottom_margin=DInches(.72); sec.left_margin=sec.right_margin=DInches(.72)
    doc.styles["Normal"].font.name="Arial"; doc.styles["Normal"].font.size=DPt(11)
    for name,size,col in [("Title",24,BLUE),("Heading 1",16,BLUE),("Heading 2",13,INK),("Heading 3",11.5,TEAL)]:
        st=doc.styles[name]; st.font.name="Arial"; st.font.size=DPt(size); st.font.bold=True; st.font.color.rgb=drgb(col)
    settings=doc.settings._element; upd=OxmlElement("w:updateFields"); upd.set(qn("w:val"),"true"); settings.append(upd)


def footer(doc:Document):
    for sec in doc.sections:
        p=sec.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        doc_font(p.add_run(f"© 2026 {ORG} · Page "),8,color=GREY)
        for instr in ["PAGE","NUMPAGES"]:
            fld=OxmlElement("w:fldSimple"); fld.set(qn("w:instr"),instr); r=OxmlElement("w:r"); t=OxmlElement("w:t"); t.text="1"; r.append(t); fld.append(r); p._p.append(fld)
            if instr=="PAGE": doc_font(p.add_run(" of "),8,color=GREY)


def cover(doc:Document,kind:str):
    logo=ASSETS/"tertiary-logo.png"
    if logo.exists():
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(logo),width=DInches(1.25))
    for text,size,bold,col,after in [(ORG,13,True,INK,0),(f"UEN: {UEN}",10,False,GREY,24),(kind.upper(),26,True,BLUE,10),("For",12,False,GREY,8),(TITLE,20,True,INK,10),(f"TGS Ref No: {CODE}",12,False,INK,18),("Conducted by",12,False,GREY,0),(ORG,13,True,INK,0),(f"UEN: {UEN}",10,False,GREY,18),(f"Version {VERSION}",12,True,BLUE,0)]:
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=DPt(after); doc_font(p.add_run(text),size,bold,col)
    doc.add_page_break()


def version_record(doc:Document,summary:str):
    p=doc.add_paragraph(); doc_font(p.add_run("DOCUMENT VERSION CONTROL RECORD"),14,True,INK)
    t=doc.add_table(rows=1,cols=4); t.style="Table Grid"; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    headers=["Version Number","Effective Date of Release","Summary of Included Changes","Author"]
    for i,h in enumerate(headers):
        cell=t.rows[0].cells[i]; cell.text=h; cell._tc.get_or_add_tcPr().append(shading(BLUE));
        for p in cell.paragraphs:
            for r in p.runs: doc_font(r,9.2,True,WHITE)
    for vals in [("2.0","26 Aug 2026","Prior AI-102 courseware release",ORG),("3.0","27 Sep 2026","AI-103 title transition",ORG),("4.0","27 Sep 2026","Five-domain AI-103 courseware and original Tertiary labs",ORG),(VERSION,DATE,summary,ORG)]:
        cells=t.add_row().cells
        for i,v in enumerate(vals): cells[i].text=v
        for cell in cells:
            for p in cell.paragraphs:
                for r in p.runs: doc_font(r,9.2)
    doc.add_page_break()


def shading(fill:str):
    shd=OxmlElement("w:shd"); shd.set(qn("w:fill"),fill); return shd


def toc(doc:Document, entries:list[str]):
    p=doc.add_paragraph(); doc_font(p.add_run("TABLE OF CONTENTS"),16,True,BLUE)
    p=doc.add_paragraph(); begin=OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"),"begin"); instr=OxmlElement("w:instrText"); instr.set(qn("xml:space"),"preserve"); instr.text=' TOC \\o "1-3" \\h \\z \\u '; sep=OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"),"separate"); end=OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"),"end")
    for el in [begin,instr,sep,end]: p._p.append(el)
    for n,entry in enumerate(entries,1):
        p=doc.add_paragraph(); p.paragraph_format.space_after=DPt(5)
        doc_font(p.add_run(f"{n:02d}  {entry}"),10.5,False,INK)
    doc.add_page_break()


def add_table(doc:Document,headers:list[str],rows:list[list[str]],width_font=9.2):
    t=doc.add_table(rows=1,cols=len(headers)); t.style="Table Grid"; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(headers):
        cell=t.rows[0].cells[i]; cell.text=h; cell._tc.get_or_add_tcPr().append(shading(BLUE))
        for p in cell.paragraphs:
            for r in p.runs: doc_font(r,width_font,True,WHITE)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,v in enumerate(row): cells[i].text=str(v)
        for cell in cells:
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri%2==0: cell._tc.get_or_add_tcPr().append(shading(LIGHT))
            for p in cell.paragraphs:
                for r in p.runs: doc_font(r,width_font)
    return t


def render_markdown(doc:Document,text:str):
    in_code=False; code=[]
    for raw in text.splitlines()[1:]:
        line=raw.rstrip()
        if line.startswith("```"):
            if in_code:
                p=doc.add_paragraph(); p._p.get_or_add_pPr().append(shading("F0F3F7")); doc_font(p.add_run("\n".join(code)),9,False,INK,"Consolas"); code=[]
            in_code=not in_code; continue
        if in_code: code.append(line); continue
        if not line.strip(): continue
        if line.startswith("### "): doc.add_paragraph(line[4:].strip(),style="Heading 3"); continue
        if line.startswith("## "): doc.add_paragraph(line[3:].strip(),style="Heading 2"); continue
        if re.match(r"^[-*]\s+",line): doc.add_paragraph(re.sub(r"^[-*]\s+","",line),style="List Bullet"); continue
        if re.match(r"^\d+\.\s+",line): doc.add_paragraph(re.sub(r"^\d+\.\s+","",line),style="List Number"); continue
        if line.startswith("|"): continue
        p=doc.add_paragraph(); doc_font(p.add_run(re.sub(r"[*`]","",line)),11)


def build_lg(labs:list[Lab]):
    d=Document(); setup_doc(d); cover(d,"Learner Guide"); version_record(d,"Pinned MicrosoftLearning exercises and assets copied unchanged; ten selected labs aligned to AI-103 domains, slides, and WSQ evidence."); toc(d,["How to Use This Guide","Before You Start"]+[f"Lab {lab.num:02d} - {lab.title}" for lab in labs]+["Quick Command Reference","Assessment Flow and Support"])
    d.add_heading("How to Use This Guide",level=1)
    d.add_paragraph("Use the trainer deck to understand mechanisms and decisions. Use this guide for the detailed click paths, commands, code, expected results, diagnostics, and evidence required in each lab.")
    add_table(d,["Authority","Current reference"],[["Approved course",f"{TITLE} · {CODE} · 2 days / 16 hours"],["Course enquiries",COURSE_URL],["AI-103 study guide","Skills measured from 16 Apr 2026"],["Microsoft course","AI-103T00-A: Develop AI apps and agents on Azure"],["LMS",LMS_URL]])
    d.add_heading("Before You Start",level=1)
    for x in ["Use the Azure subscription or lab environment supplied by the trainer.","Install current Azure CLI/Python tooling only when the lab requires it.","Authenticate with Microsoft Entra ID where supported; never save live keys in the repository.","Use synthetic data and remove training resources after evidence has been captured.","For every lab, retain request IDs, deployment/model versions, screenshots or JSON outputs, and a short interpretation."]:
        d.add_paragraph(x,style="List Bullet")
    for lab in labs:
        d.add_heading(f"Lab {lab.num:02d} - {lab.title}",level=1)
        d.add_paragraph(f"Source folder: {lab.path.parent.relative_to(ROOT) if lab.path.name=='README.md' else lab.path.relative_to(ROOT)}")
        render_markdown(d,lab.text)
        official_path = LABS / OFFICIAL[lab.num-1]["source_path"]
        official_text = official_path.read_text(encoding="utf-8")
        official_text = re.sub(r"\A---\s*\n.*?\n---\s*\n", "", official_text, flags=re.S)
        d.add_heading("Unchanged MicrosoftLearning exercise instructions",level=2)
        d.add_paragraph(f"Exact Markdown source and companion assets: labs/{OFFICIAL[lab.num-1]['source_path']}. Follow the numbered instructions below with the source assets. MicrosoftLearning MIT license and pinned revision are in labs/official/ and the source manifest.")
        render_markdown(d,official_text)
        d.add_heading("Acceptance Evidence",level=2)
        d.add_paragraph(TOPICS[lab.num-1]["verify"])
        d.add_heading("Troubleshooting Checklist",level=2)
        for x in ["Capture the full error and request/operation ID before retrying.","Confirm the endpoint, API/model version, region, identity, and RBAC scope.","Validate payload/schema fields and input content type.","Check quota, rate limits, network resolution, private connectivity, and service status.","Record the corrective action and rerun the acceptance check."]:
            d.add_paragraph(x,style="List Number")
    d.add_heading("Quick Command Reference",level=1)
    add_table(d,["Command / item","Purpose"],[["az login --use-device-code","Authenticate Azure CLI without embedding credentials"],["az account show --output table","Confirm active tenant/subscription"],["python -m venv .venv","Create an isolated Python environment"],["az monitor metrics list ...","Retrieve a metric for operational evidence"],["git status --short","Confirm no secrets or generated private assessment files are staged"]])
    d.add_heading("Assessment Flow and Support",level=1)
    for x in ["Complete TRAQOM digital attendance.","Complete Assessment digital attendance.","Sit the 60-minute WA-SAQ, then the 60-minute PP.","Submit the two candidate papers on the LMS.","Sign the Assessment Summary Record."]:
        d.add_paragraph(x,style="List Number")
    d.add_paragraph(f"Support: enquiry@tertiaryinfotech.com · +65 6100 0613 · {LMS_URL}")
    footer(d); d.save(LG)


def build_lp(labs:list[Lab],m:dict[str,int]):
    d=Document(); setup_doc(d); cover(d,"Lesson Plan"); version_record(d,"Updated to v5.0 deck slide map, pinned MicrosoftLearning lab sources, five AI-103 domains, and approved 60+60 minute assessments."); toc(d,["Course Overview","Learning Outcomes","Daily Schedule","Topic-by-Topic Breakdown","Resources Required","Assessment"])
    d.add_heading("Course Overview",level=1)
    d.add_paragraph("A two-day instructor-led WSQ course that develops the ability to analyse, evaluate, and improve Azure AI solutions through Microsoft Foundry services, governed engineering decisions, and observable evidence.")
    d.add_heading("Learning Outcomes",level=1)
    for x in ["LO1: Analyse algorithms in Microsoft Azure AI applications and establish their correlation with efficiency.","LO2: Identify and evaluate strengths and limitations of Microsoft Azure AI applications.","LO3: Assess feasibility and improvements when applying Azure AI to product and maintenance processes."]:
        d.add_paragraph(x,style="List Bullet")
    d.add_heading("Daily Schedule",level=1)
    rows=[]
    rows += [["Day 1 · 09:00–09:40","40 min","Administration, outcomes, assessment briefing","Facilitated briefing",f"Slides 1–{m['admin_labs']}"]]
    slots1=[("09:40–10:45",1),("11:00–12:05",2),("12:05–13:00",3),("14:00–15:05",4),("15:05–16:10",5),("16:25–17:35",6)]
    for tm,n in slots1: rows.append([f"Day 1 · {tm}","55–70 min",f"Topic/Lab {n:02d}: {labs[n-1].title}","Mechanism + evidence lab",f"Slides {m[f'topic{n:02d}']}–{m[f'lab{n:02d}_verify']}"])
    rows += [["Day 1 · 10:45–11:00","15 min","Morning break","Break","—"],["Day 1 · 13:00–14:00","60 min","Lunch","Break","—"],["Day 1 · 16:10–16:25","15 min","Afternoon break","Break","—"],["Day 1 · 17:35–18:00","25 min","Evidence review and recap","Discussion",f"Slides {m['topic01']}–{m['lab06_verify']}"]]
    rows += [["Day 2 · 09:00–09:15","15 min","Digital attendance and recap","Review",f"Slide {m['day2']}"]]
    slots2=[("09:15–10:20",7),("10:35–11:45",8),("11:45–13:00",9),("14:00–15:15",10)]
    for tm,n in slots2: rows.append([f"Day 2 · {tm}","65–75 min",f"Topic/Lab {n:02d}: {labs[n-1].title}","Mechanism + evidence lab",f"Slides {m[f'topic{n:02d}']}–{m[f'lab{n:02d}_verify']}"])
    rows += [["Day 2 · 10:20–10:35","15 min","Morning break","Break","—"],["Day 2 · 13:00–14:00","60 min","Lunch","Break","—"],["Day 2 · 15:15–15:45","30 min","Assessment briefing and preparation","Briefing",f"Slides {m['admin_briefing']}–{m['admin_assessment_flow']}"] ,["Day 2 · 16:00–18:00","120 min","WA-SAQ (60 min) + PP (60 min)","Open-book assessment",f"Slides {m['closing_assessment']}–{m['closing_attendance']}"]]
    add_table(d,["Time","Duration","Topic / Activity","Method","Slides"],rows,8.3)
    d.add_heading("Topic-by-Topic Breakdown",level=1)
    for n,(lab,t) in enumerate(zip(labs,TOPICS),1):
        d.add_heading(f"Topic {n:02d} - {t['title']} · Slides {m[f'topic{n:02d}']}–{m[f'lab{n:02d}_verify']}",level=2)
        d.add_paragraph(f"Lab {n:02d} - {lab.title} · Slide {m[f'lab{n:02d}_activity']}",style="Heading 3")
        d.add_paragraph("Coverage: architecture and boundaries; runtime mechanism; exact fields/endpoints; implementation excerpt; editable metric; decision rule; failure/control; activity bridge; acceptance evidence.")
    d.add_heading("Resources Required",level=1)
    for x in ["Trainer PPTX and learner PDF", "Learner Guide", "Azure subscription or trainer-provided lab environment", "Modern browser, Azure CLI, Python 3, and Visual Studio Code as required", "LMS/TMS access for courseware and assessment"]: d.add_paragraph(x,style="List Bullet")
    d.add_heading("Assessment",level=1)
    d.add_paragraph("Approved instruments: Written Assessment - Short-Answer Questions (WA-SAQ), 60 minutes, K1–K6; Practical Performance (PP), 60 minutes, A1–A6. Both are individual and open book. Complete assessment digital attendance before starting and submit candidate papers on the LMS.")
    footer(d); d.save(LP)


def main():
    CW.mkdir(exist_ok=True); ASSETS.mkdir(exist_ok=True)
    labs=load_labs(); deck=Deck(labs); slide_map=deck.build(); build_lg(labs); build_lp(labs,slide_map)
    print(f"Wrote {PPTX} ({len(deck.prs.slides)} slides)")
    print(f"Wrote {LG}")
    print(f"Wrote {LP}")
    print(f"Wrote {SLIDE_MAP}")


if __name__ == "__main__":
    main()
